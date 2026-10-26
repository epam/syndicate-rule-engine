import json
import re
from collections import defaultdict
from collections.abc import Iterable, Iterator
from copy import copy
from datetime import UTC, datetime
from io import BytesIO
from typing import BinaryIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from modular_sdk.models.tenant import Tenant
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.packaging.relationship import get_dependents, get_rels_path
from openpyxl.utils.cell import get_column_letter, range_boundaries

from helpers.constants import Cloud, SettingKey
from helpers.reports import service_from_resource_type
from models.job import Job
from models.resource import Resource
from services.clients.s3 import S3Client
from services.environment_service import EnvironmentService
from services.license_service import LicenseService
from services.metadata import Metadata, RuleMetadata
from services.platform_service import PlatformService
from services.report_service import StatisticsItem
from services.reports_bucket import (
    PlatformReportsBucketKeysBuilder,
    ReportsBucketKeysBuilder,
    TenantReportsBucketKeysBuilder,
)
from services.resources import (
    CloudResource,
    K8SResource,
    iter_rule_region_resources,
    prepare_resource_type,
)
from services.resources_service import ResourcesService
from services.setting_service import SettingsService
from services.sharding import RuleMeta, ShardsCollection


TEMPLATE_KEY = ReportsBucketKeysBuilder.dashboard_template_key()
EMPTY_RULE_META: RuleMeta = {'description': '', 'resource': '', 'comment': ''}
SEVERITY_PRIORITIES = {
    'High': 40,
    'Medium': 30,
    'Low': 20,
    'Info': 10,
}
REMEDIATION_COMPLEXITY_PRIORITIES = {
    'Low': 5,
    'Low-Medium': 4,
    'Medium': 3,
    'Medium-High': 2,
    'High': 1,
}
RAW_DATA_COLUMNS = {
    'ResourceType': 'resource_type',
    'RuleID': 'rule_id',
    'RuleTitle': 'rule_title',
    'ResourceID': 'resource_id',
    'ResourceTags': 'resource_tags',
    'CheckStatus': 'check_status',
    'Severity': 'severity',
    'SeverityPriority': 'severity_priority',
    'RemediationComplexity': 'remediation_complexity',
    'RemediationComplexityPriority': 'remediation_complexity_priority',
    'RemediationPriority': 'remediation_priority',
    'Region (Namespace)': 'region_namespace',
    'Date': 'date',
    'Service': 'service',
    'ServiceCategory': 'service_section',
    'Category': 'category',
}
SHORT_RULE_ID_REGEX = re.compile(r'[^-]+-\d+(?=-|$)')

_TABLE_NAME = 'raw_table'
_MAX_ROWS = 1_048_576
_MAX_COLUMNS = 16_384


def _locate_table(workbook):
    matches = [
        (sheet, table)
        for sheet in workbook.worksheets
        for table in sheet.tables.values()
        if table.name == _TABLE_NAME
    ]
    if len(matches) != 1:
        raise ValueError(
            f'Template must contain exactly one {_TABLE_NAME!r} table'
        )
    return matches[0]


def _request_formula_recalculation(workbook) -> None:
    calculation = workbook.calculation
    calculation.calcMode = 'auto'
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True
    calculation.calcOnSave = True


def _template_chart_paths(archive: ZipFile) -> dict[str, list[str]]:
    relationships = {
        rel.Id: rel
        for rel in get_dependents(archive, 'xl/_rels/workbook.xml.rels')
    }
    root = ElementTree.fromstring(archive.read('xl/workbook.xml'))
    relationship_id = (
        '{http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships}id'
    )
    paths = {}
    for sheet in root.findall('./{*}sheets/{*}sheet'):
        sheet_path = relationships[sheet.attrib[relationship_id]].target
        sheet_root = ElementTree.fromstring(archive.read(sheet_path))
        paths[sheet.attrib['name']] = charts = []
        drawings = sheet_root.findall('./{*}drawing')
        if not drawings:
            continue
        sheet_rels = {
            rel.Id: rel
            for rel in get_dependents(archive, get_rels_path(sheet_path))
        }
        for drawing in drawings:
            drawing_path = sheet_rels[drawing.attrib[relationship_id]].target
            drawing_model = SpreadsheetDrawing.from_tree(
                ElementTree.fromstring(archive.read(drawing_path))
            )
            chart_rels = drawing_model._chart_rels
            if not chart_rels:
                continue
            drawing_rels = {
                rel.Id: rel
                for rel in get_dependents(archive, get_rels_path(drawing_path))
            }
            charts.extend(drawing_rels[rel.id].target for rel in chart_rels)
    return paths


def _restore_template_charts(
    template: BinaryIO, generated: BytesIO, workbook
) -> BytesIO:
    output = BytesIO()
    with (
        ZipFile(template) as original,
        ZipFile(generated) as updated,
        ZipFile(output, 'w', compression=ZIP_DEFLATED) as restored,
    ):
        chart_paths = _template_chart_paths(original)
        replacements = {}
        pending = []
        for sheet in workbook.worksheets:
            sources = chart_paths.get(sheet.title, [])
            if len(sources) != len(sheet._charts):
                raise ValueError(
                    f'Cannot preserve all charts on worksheet {sheet.title!r}'
                )
            pending.extend(
                (source, chart.path.lstrip('/'))
                for source, chart in zip(sources, sheet._charts)
            )

        original_types = ElementTree.fromstring(
            original.read('[Content_Types].xml')
        )
        content_types = ElementTree.fromstring(
            updated.read('[Content_Types].xml')
        )
        overrides = {
            node.attrib['PartName'].lstrip('/'): node.attrib['ContentType']
            for node in original_types.findall('./{*}Override')
        }
        defaults = {
            node.attrib['Extension']: node.attrib['ContentType']
            for node in original_types.findall('./{*}Default')
        }
        type_tag = original_types.tag.replace('Types', 'Override')
        while pending:
            source, destination = pending.pop()
            if destination in replacements:
                continue
            # Keep chart XML intact; namespace dependencies to avoid collisions
            # with media and other parts renumbered by openpyxl.
            replacements[destination] = original.read(source)
            for node in list(content_types):
                if node.attrib.get('PartName') == '/' + destination:
                    content_types.remove(node)
            content_type = overrides.get(source)
            if content_type is None:
                content_type = defaults[source.rsplit('.', 1)[-1]]
            ElementTree.SubElement(
                content_types,
                type_tag,
                PartName='/' + destination,
                ContentType=content_type,
            )
            rels_path = get_rels_path(source)
            if rels_path not in original.namelist():
                continue
            relationships = {
                rel.Id: rel for rel in get_dependents(original, rels_path)
            }
            rels_root = ElementTree.fromstring(original.read(rels_path))
            for rel in rels_root:
                if rel.attrib.get('TargetMode') == 'External':
                    continue
                target = relationships[rel.attrib['Id']].target
                preserved_target = 'xl/preservedChartParts/' + target
                rel.set('Target', '/' + preserved_target)
                pending.append((target, preserved_target))
            replacements[get_rels_path(destination)] = ElementTree.tostring(
                rels_root
            )

        replacements['[Content_Types].xml'] = ElementTree.tostring(
            content_types
        )
        for entry in updated.infolist():
            restored.writestr(
                entry,
                replacements.pop(entry.filename, updated.read(entry)),
            )
        for path, data in replacements.items():
            restored.writestr(path, data)
    output.seek(0)
    return output


def fill_dashboard_template(
    template: BinaryIO,
    rows: list[dict],
    job_details: dict[str, str] | None = None,
) -> BytesIO:
    template = BytesIO(template.read())
    workbook = load_workbook(template)
    sheet, table = _locate_table(workbook)
    min_col, min_row, max_col, old_max_row = range_boundaries(table.ref)
    columns = [column.name for column in table.tableColumns]

    new_max_row = min_row + len(rows)
    if new_max_row > _MAX_ROWS or max_col > _MAX_COLUMNS:
        raise ValueError('Dashboard data exceeds Excel worksheet limits')

    if new_max_row > old_max_row:
        for row_num in range(old_max_row + 1, new_max_row + 1):
            if any(
                (
                    (cell := sheet.cell(row_num, col_num)).value is not None
                    or cell.comment is not None
                    or cell.hyperlink is not None
                )
                for col_num in range(min_col, max_col + 1)
            ):
                raise ValueError(
                    'Expanding raw_table would overwrite unrelated '
                    f'worksheet data at row {row_num}'
                )
        for merged_range in sheet.merged_cells.ranges:
            if (
                merged_range.min_row <= new_max_row
                and merged_range.max_row > old_max_row
                and merged_range.min_col <= max_col
                and merged_range.max_col >= min_col
            ):
                raise ValueError(
                    'Expanding raw_table would overlap merged worksheet cells'
                )

        style_row = max(min_row + 1, old_max_row)
        for row_num in range(old_max_row + 1, new_max_row + 1):
            for col_num in range(min_col, max_col + 1):
                sheet.cell(row_num, col_num)._style = copy(
                    sheet.cell(style_row, col_num)._style
                )

    for row_num in range(min_row + 1, old_max_row + 1):
        for col_num in range(min_col, max_col + 1):
            sheet.cell(row_num, col_num).value = None

    for row_offset, row in enumerate(rows, start=1):
        for col_offset, column in enumerate(columns):
            cell = sheet.cell(min_row + row_offset, min_col + col_offset)
            data_key = RAW_DATA_COLUMNS.get(column)
            value = row.get(data_key, '') if data_key else ''
            if isinstance(value, str):
                cell.value = value
                cell.data_type = 's'
            else:
                cell.value = value

    ref = (
        f'{get_column_letter(min_col)}{min_row}:'
        f'{get_column_letter(max_col)}{new_max_row}'
    )
    table.ref = ref
    table.autoFilter.ref = ref

    if job_details is not None:
        introduction = workbook['Introduction']
        for row_number, label in enumerate(
            (
                'Job ID',
                'Tenant name',
                'Platform',
                'Submitted at',
                'Started at',
                'Finished at',
                'Account ID',
            ),
            start=3,
        ):
            if introduction.cell(row_number, 2).value != label:
                raise ValueError(
                    f'Introduction!B{row_number} must contain {label!r}'
                )
            if label == 'Account ID' and 'Platform ID' in job_details:
                label = 'Platform ID'
                introduction.cell(row_number, 2, label)
            value_cell = introduction.cell(row_number, 3, job_details[label])
            value_cell.data_type = 's'

    _request_formula_recalculation(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    template.seek(0)
    return _restore_template_charts(template, output, workbook)


def resolve_resource_type_name(
    resource_type: str,
    cloud: Cloud,
    resource_names_map: dict,
) -> str:
    if not resource_type:
        return ''

    normalized_type = prepare_resource_type(resource_type, cloud)
    cloud_map = resource_names_map.get(cloud.value)
    if cloud_map is None:
        return service_from_resource_type(normalized_type)

    if name := cloud_map.get(normalized_type):
        return name

    return service_from_resource_type(normalized_type)


class DashboardReportError(Exception):
    def __init__(self, message: str, status_code: int):
        super().__init__(message)
        self.status_code = status_code


def _iter_tenant_inventory(
    resources_service: ResourcesService,
    tenant_name: str,
    customer_name: str,
    platform_id: str | None = None,
) -> Iterator[Resource]:

    page = resources_service.get_resources(
        tenant_name=tenant_name,
        customer_name=customer_name,
        platform_id=platform_id,
        limit=50,
    )
    yield from page

    while page.last_evaluated_key:
        page = resources_service.get_resources(
            tenant_name=tenant_name,
            customer_name=customer_name,
            platform_id=platform_id,
            limit=50,
            last_evaluated_key=page.last_evaluated_key,
        )
        yield from page


def _row(
    rule_id: str,
    resource: CloudResource | None,
    region: str,
    status: str,
    metadata: RuleMetadata,
    scan_metadata: RuleMeta,
    cloud: Cloud,
) -> dict:
    short_rule_id = rule_id.removeprefix('ecc-')
    numbered_rule_id = SHORT_RULE_ID_REGEX.match(short_rule_id)
    if numbered_rule_id:
        short_rule_id = numbered_rule_id.group()

    report_region = region
    if resource and isinstance(resource, K8SResource):
        report_region = resource.namespace or ''

    if cloud is Cloud.KUBERNETES and resource is None:
        report_region = ''

    severity = metadata.severity.value if metadata.severity else ''
    remediation_complexity = (
        metadata.remediation_complexity.value
        if metadata.remediation_complexity
        else ''
    )
    severity_priority = SEVERITY_PRIORITIES.get(severity, 0)
    remediation_complexity_priority = REMEDIATION_COMPLEXITY_PRIORITIES.get(
        remediation_complexity, 0
    )
    return {
        'resource_type': (
            prepare_resource_type(resource.resource_type, cloud)
            if resource
            else prepare_resource_type(scan_metadata['resource'], cloud)
        ),
        'rule_id': short_rule_id,
        'rule_title': scan_metadata['description'],
        'resource_id': resource.id if resource else '',
        'resource_tags': json.dumps(
            resource.tags if resource else {},
            ensure_ascii=False,
            sort_keys=True,
        ),
        'check_status': status,
        'severity': severity,
        'severity_priority': severity_priority,
        'remediation_complexity': remediation_complexity,
        'remediation_complexity_priority': remediation_complexity_priority,
        'remediation_priority': (
            severity_priority + remediation_complexity_priority
        ),
        'region_namespace': report_region,
        'date': (
            datetime.fromtimestamp(resource.sync_date, tz=UTC).isoformat()
            if resource
            else ''
        ),
        'service': metadata.service,
        'service_section': metadata.service_section,
        'category': '.'.join(
            part.strip() for part in metadata.category.split('>')[:2]
        ),
    }


def _register_rule_region(
    pair: tuple[str, str],
    scan_metadata: dict[str, RuleMeta],
    cloud: Cloud,
    candidates_by_region_type: dict[tuple[str, str], set[tuple[str, str]]],
    indexed_rule_regions: set[tuple[str, str]],
) -> None:
    if pair in indexed_rule_regions:
        return
    indexed_rule_regions.add(pair)
    rule_entry = scan_metadata.get(pair[0], EMPTY_RULE_META)
    if rule_entry['resource']:
        resource_type = prepare_resource_type(rule_entry['resource'], cloud)
        candidates_by_region_type[(pair[1], resource_type)].add(pair)


def make_dashboard_rows(
    findings: Iterable[tuple[str, str, Iterable[CloudResource]]],
    statistics: Iterable[StatisticsItem],
    inventory: Iterable[Resource],
    metadata: Metadata,
    cloud: Cloud,
    scan_metadata: dict[str, RuleMeta] | None = None,
) -> list[dict]:
    """Reconcile findings and inventory, preserving inventory row details."""
    scan_metadata = scan_metadata or {}
    normalize = str.casefold if cloud is Cloud.AZURE else str
    candidates_by_region_type = defaultdict(set)
    indexed_rule_regions = set()

    evaluation_statuses = {}
    for item in statistics:
        pair = (item.policy, item.region)
        evaluation_statuses[pair] = (
            'Passed' if item.is_successful() else 'Not evaluated'
        )
        _register_rule_region(
            pair,
            scan_metadata,
            cloud,
            candidates_by_region_type,
            indexed_rule_regions,
        )

    violations = {}
    for rule_id, region, resources in findings:
        pair = (rule_id, region)
        _register_rule_region(
            pair,
            scan_metadata,
            cloud,
            candidates_by_region_type,
            indexed_rule_regions,
        )
        for resource in resources:
            resource_type = prepare_resource_type(
                resource.resource_type, cloud
            )
            candidates_by_region_type[(region, resource_type)].add(pair)
            violations.setdefault(
                (pair, resource_type, normalize(resource.id)), resource
            )

    candidates = {}
    for inventory_resource in inventory:
        resource = inventory_resource.to_cloud_resource()
        region = str(resource.location)
        resource_type = prepare_resource_type(resource.resource_type, cloud)
        resource_id = normalize(resource.id)

        for pair in candidates_by_region_type[(region, resource_type)]:
            if (pair, resource_type, resource_id) in violations:
                status = 'Failed'
            elif pair in evaluation_statuses:
                status = evaluation_statuses[pair]
            else:
                continue
            candidates.setdefault(
                (pair, resource_type, resource_id), (resource, status)
            )

    for (pair, resource_type, resource_id), resource in violations.items():
        candidates.setdefault(
            (pair, resource_type if resource_id else '', resource_id),
            (resource, 'Failed'),
        )

    represented_rule_regions = {pair for pair, _, _ in candidates}
    for pair in sorted(
        pair
        for pair, status in evaluation_statuses.items()
        if status == 'Not evaluated' and pair not in represented_rule_regions
    ):
        candidates[(pair, '', '')] = (None, 'Not evaluated')

    output = []
    rules_metadata = {}
    for (
        ((rule_id, region), _, _),
        (resource, status),
    ) in candidates.items():
        entry = scan_metadata.get(rule_id, EMPTY_RULE_META)
        # Without a comment, Metadata.rule's resource-only fallback cannot
        # resolve the rule's cloud. Use the stored metadata in that case.
        resource_type = None
        if entry['comment']:
            resource_type = (
                prepare_resource_type(resource.resource_type, cloud)
                if resource
                else entry['resource']
            )

        metadata_key = (rule_id, resource_type)
        if metadata_key not in rules_metadata:
            rules_metadata[metadata_key] = metadata.rule(
                rule_id,
                comment=entry['comment'] or None,
                resource=resource_type,
            )

        output.append(
            _row(
                rule_id=rule_id,
                resource=resource,
                region=region,
                status=status,
                metadata=rules_metadata[metadata_key],
                scan_metadata=entry,
                cloud=cloud,
            )
        )

    return sorted(
        output,
        key=lambda row: (
            row['resource_type'],
            row['rule_id'],
            row['region_namespace'],
            row['resource_id'],
        ),
    )


class DashboardReportService:
    def __init__(
        self,
        s3_client: S3Client,
        environment_service: EnvironmentService,
        license_service: LicenseService,
        resources_service: ResourcesService,
        platform_service: PlatformService,
        settings_service: SettingsService,
    ):
        self._s3_client = s3_client
        self._environment_service = environment_service
        self._license_service = license_service
        self._resources_service = resources_service
        self._platform_service = platform_service
        self._settings_service = settings_service

    def collect_raw_data(
        self,
        job: Job,
        tenant: Tenant,
        collection: ShardsCollection,
        statistics: Iterable[StatisticsItem],
        keys_builder: ReportsBucketKeysBuilder,
    ) -> None:
        """Persist Excel rows before the scan job is marked completed."""
        cloud = keys_builder.cloud
        inventory = _iter_tenant_inventory(
            self._resources_service,
            tenant.name,
            tenant.customer_name,
            job.platform_id if job.is_platform_job else None,
        )

        metadata = self._license_service.get_customer_metadata(
            tenant.customer_name
        )

        rows = make_dashboard_rows(
            findings=iter_rule_region_resources(collection, cloud, metadata),
            statistics=statistics,
            inventory=inventory,
            metadata=metadata,
            cloud=cloud,
            scan_metadata=collection.meta,
        )

        self._s3_client.gz_put_json(
            bucket=(
                self._environment_service.default_reports_bucket_name()
            ),
            key=keys_builder.job_dashboard_raw(job),
            obj=rows,
        )

    def generate(self, job: Job, tenant: Tenant) -> BytesIO:
        if not job.is_succeeded or not job.is_finished:
            raise DashboardReportError(
                'Dashboard reports require a completed successful job', 400
            )
        if (
            job.tenant_name != tenant.name
            or job.customer_name != tenant.customer_name
        ):
            raise DashboardReportError(
                'The requested job does not belong to this tenant', 400
            )

        keys_builder: ReportsBucketKeysBuilder
        if job.is_platform_job:
            platform = self._platform_service.get_nullable(
                hash_key=job.platform_id
            )
            if not platform:
                raise DashboardReportError(
                    f'Platform {job.platform_id!r} not found for this tenant',
                    404,
                )
            keys_builder = PlatformReportsBucketKeysBuilder(platform)
        else:
            keys_builder = TenantReportsBucketKeysBuilder(tenant)

        bucket = self._environment_service.default_reports_bucket_name()

        try:
            rows = self._s3_client.gz_get_json(
                bucket=bucket,
                key=keys_builder.job_dashboard_raw(job),
            )
        except Exception as error:
            raise DashboardReportError(
                f'Dashboard raw data is invalid for job {job.id}', 400
            ) from error

        if not rows:
            raise DashboardReportError(
                f'Dashboard raw data is unavailable for job {job.id}', 404
            )

        resource_names_map = self._settings_service.get(
            SettingKey.CUSTODIAN_RESOURCE_NAMES_MAP
        ) or {}

        for row in rows:
            row['resource_type'] = resolve_resource_type_name(
                    row['resource_type'], keys_builder.cloud, resource_names_map
                )

        template = self._s3_client.get_object(
            bucket=bucket,
            key=TEMPLATE_KEY,
        )
        if template is None:
            raise DashboardReportError(
                f'Dashboard workbook template is missing at {TEMPLATE_KEY}',
                503,
            )

        job_details = {
            'Job ID': job.id,
            'Tenant name': tenant.display_name or tenant.name,
            'Platform': keys_builder.cloud.value,
            'Submitted at': job.submitted_at or 'Not recorded',
            'Started at': job.started_at or 'Not recorded',
            'Finished at': job.stopped_at or 'Not recorded',
        }

        if job.is_platform_job:
            job_details['Platform ID'] = platform.description
        else:
            job_details['Account ID'] = tenant.project or 'Not recorded'

        try:
            workbook = fill_dashboard_template(
                template, rows, job_details=job_details
            )
        except Exception as error:
            raise DashboardReportError(
                f'Incompatible dashboard workbook template: {error}', 503
            ) from error

        return workbook
